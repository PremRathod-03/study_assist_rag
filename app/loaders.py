import pymupdf4llm
from pptx import Presentation
from docx import Document


def load_pdf(path: str) -> list[dict]:
    """
    Reads a PDF file and returns its content as a list of pages.
    Uses pymupdf4llm instead of a plain PDF reader because it converts
    each page to Markdown, preserving headings/bullets/tables structure
    instead of dumping one long unstructured blob of text.
    """
    # page_chunks=True makes it return one dict per page instead of
    # one giant string for the whole document
    pages = pymupdf4llm.to_markdown(path, page_chunks=True)

    # Reshape into our own consistent format: every loader in this file
    # returns the same shape {"text": ..., "source_type": ..., page/slide}
    # so nothing downstream needs to know or care what the original file was
    return [
        {"text": pg["text"], "source_type": "pdf_page", "page": i}
        for i, pg in enumerate(pages)
    ]


def load_pptx(path: str) -> list[dict]:
    """
    Reads a PowerPoint file and returns its content as a list of slides.
    Pulls text from every text box on the slide, AND from the speaker
    notes (since notes often contain the actual explanation, while the
    slide itself might just have a diagram or a few bullet points).
    """
    prs = Presentation(path)
    slides = []

    for i, slide in enumerate(prs.slides):
        # Collect text from every shape (textbox, title, etc.) on this slide
        text_parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text_parts.append(shape.text_frame.text)

        # Also grab the speaker notes, if this slide has any
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text

        # Combine slide text + notes into one block for this slide
        combined = "\n".join(text_parts)
        if notes:
            combined += "\n[Notes] " + notes

        slides.append({"text": combined, "source_type": "pptx_slide", "slide": i})

    return slides


def load_docx(path: str) -> list[dict]:
    """
    Reads a Word document and returns its content.
    Simpler than PDF/PPTX because docx text is already just paragraphs,
    no layout guessing needed.
    """
    doc = Document(path)
    full_text = "\n".join(p.text for p in doc.paragraphs)

    # Returned as a list (with one item) to match the same shape as
    # load_pdf and load_pptx, even though there's only one "section" here
    return [{"text": full_text, "source_type": "docx_section", "page": 0}]


def load_file(path: str) -> list[dict]:
    """
    The single entry point other files should actually call.
    Looks at the file extension and routes to the right loader above,
    so callers never need to know or check the file type themselves.
    """
    if path.endswith(".pdf"):
        return load_pdf(path)
    if path.endswith(".pptx"):
        return load_pptx(path)
    if path.endswith(".docx"):
        return load_docx(path)
    raise ValueError(f"Unsupported file type: {path}")